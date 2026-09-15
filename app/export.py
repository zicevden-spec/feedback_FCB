from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app import db, roles


def _status_title(rec):
    if rec["published"]:
        return "Отвечено в чате"
    if rec["answer_text"]:
        return "Ответ написан, не опубликован"
    return "Без ответа"


def _autowidth(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 60)
    for c in ws[1]:
        c.font = Font(bold=True)


def build_report_xlsx() -> bytes:
    wb = Workbook()

    # ---------- Статистика ----------
    ws = wb.active
    ws.title = "Статистика"
    s = db.stats()
    q_wait = s["q_total"] - s["q_published"] - s["q_in_progress"]
    p_wait = s["p_total"] - s["p_published"] - s["p_in_progress"]
    rows = [
        ["Отчёт чат-бота ФЦБ", ""],
        ["Сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")],
        ["", ""],
        ["Вопросы юристу", "Кол-во"],
        ["Всего", s["q_total"]],
        ["Отвечено в чате", s["q_published"]],
        ["Ответ написан, не опубликован", s["q_in_progress"]],
        ["Без ответа", q_wait],
        ["", ""],
        ["Агентские выплаты", "Кол-во"],
        ["Всего", s["p_total"]],
        ["Отвечено в чате", s["p_published"]],
        ["Ответ написан, не опубликован", s["p_in_progress"]],
        ["Без ответа", p_wait],
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 20

    # ---------- Вопросы ----------
    ws = wb.create_sheet("Вопросы")
    ws.append(["ID", "Дата", "Клиент", "ФИО", "Город", "Телефон", "Вопрос", "Статус"])
    with db.get_conn() as conn:
        items = conn.execute("SELECT * FROM questions ORDER BY id").fetchall()
    for r in items:
        ws.append([r["id"], r["created_at"], r["client_mention"], r["fio"], r["city"], r["phone"], r["question_text"], _status_title(r)])
    _autowidth(ws)

    # ---------- Выплаты ----------
    ws = wb.create_sheet("Выплаты")
    ws.append(["ID", "Дата", "Агент", "ФИО агента", "Телефон агента", "Клиент", "ФИО клиента", "Телефон клиента", "Вопрос", "Статус"])
    with db.get_conn() as conn:
        items = conn.execute("SELECT * FROM payouts ORDER BY id").fetchall()
    for r in items:
        ws.append([r["id"], r["created_at"], r["client_mention"], r["agent_fio"], r["agent_phone"], r["client_fio"], r["client_phone"], r["question_text"], _status_title(r)])
    _autowidth(ws)

    # ---------- Сотрудники ----------
    ws = wb.create_sheet("Сотрудники")
    ws.append(["Telegram ID", "Роль"])
    for srow in db.list_staff():
        ws.append([srow["user_id"], roles.ROLE_TITLES.get(srow["role"], srow["role"])])
    _autowidth(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
